import unittest
from pathlib import Path
from backend.pipeline_context import PipelineContext
from backend.pipeline.stages import s01_document_processing


class TestDocumentInput(unittest.TestCase):
    def test_text_upload(self):
        text = 'WORK EXPERIENCE\nDeveloper at Example Ltd January 2020 to December 2022\n' * 4
        ctx = PipelineContext('test', 'resume.txt', raw_bytes=text.encode('utf-8-sig'))
        s01_document_processing(ctx)
        self.assertIn('Developer at Example', ctx.raw_text)
        self.assertNotIn('\ufeff', ctx.raw_text)

    def test_pdf_download_names(self):
        raw = Path('dataset/resumes/GopalPrasad K.pdf').read_bytes()
        for name in ('resume.pdf?token=123', 'download'):
            with self.subTest(name=name):
                ctx = PipelineContext('test', name, raw_bytes=raw)
                result = s01_document_processing(ctx)
                self.assertNotEqual(result.status, 'FAILED')
                self.assertGreater(len(ctx.raw_text), 100)

    def test_docx_table_stays_under_its_original_section(self):
        import io
        from docx import Document
        doc = Document()
        doc.add_paragraph('WORK EXPERIENCE')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Developer at Alpha Ltd'
        table.cell(0, 1).text = 'Jan 2020 - Dec 2022'
        doc.add_paragraph('EDUCATION')
        doc.add_paragraph('B.Tech, University, 2015 - 2019')
        data = io.BytesIO()
        doc.save(data)
        ctx = PipelineContext('table', 'resume.docx', raw_bytes=data.getvalue())
        s01_document_processing(ctx)
        self.assertLess(ctx.raw_text.index('WORK EXPERIENCE'), ctx.raw_text.index('Alpha Ltd'))
        self.assertLess(ctx.raw_text.index('Alpha Ltd'), ctx.raw_text.index('EDUCATION'))

    def test_missing_pdf_dependency_explains_environment_fix(self):
        import builtins
        from unittest.mock import patch
        real_import = builtins.__import__
        def without_pdfplumber(name, *args, **kwargs):
            if name == 'pdfplumber':
                raise ModuleNotFoundError("No module named 'pdfplumber'", name='pdfplumber')
            return real_import(name, *args, **kwargs)
        ctx = PipelineContext('missing-dep', 'resume.pdf', raw_bytes=b'%PDF-test')
        with patch('builtins.__import__', side_effect=without_pdfplumber):
            result = s01_document_processing(ctx)
        self.assertEqual(result.status, 'FAILED')
        self.assertIn("missing Python dependency 'pdfplumber'", result.errors[0])
        self.assertIn('requirements.txt', result.errors[0])
        self.assertIn('interpreter:', result.errors[0])

    def test_broken_docx_relationship_recovers_original_body_conservatively(self):
        import io
        import zipfile
        from docx import Document
        from backend.pipeline import runner
        doc=Document()
        doc.add_paragraph('WORK EXPERIENCE')
        doc.add_paragraph('Engineer, Alpha Ltd, Jan 2020 - Dec 2021')
        doc.add_paragraph('Engineer, Beta Ltd, Jan 2023 - Dec 2024')
        original=io.BytesIO();doc.save(original)
        damaged=io.BytesIO()
        with zipfile.ZipFile(original) as source_zip, zipfile.ZipFile(damaged,'w') as target:
            for name in source_zip.namelist():
                data=source_zip.read(name)
                if name=='word/_rels/document.xml.rels':
                    data=data.replace(b'</Relationships>',b'<Relationship Id="rId999" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="NULL"/></Relationships>')
                target.writestr(name,data)
        ctx=PipelineContext('damaged','damaged.docx',raw_bytes=damaged.getvalue())
        runner.run(ctx)
        self.assertIn('Engineer, Alpha Ltd, Jan 2020 - Dec 2021',ctx.raw_text)
        self.assertEqual(ctx.recruiter_output['status'],'PARTIAL')
        self.assertTrue(any('broken package' in w for w in ctx.stage_results['document_processing'].warnings))
        self.assertEqual([g.state for g in ctx.gaps],['INSUFFICIENT_EVIDENCE'])
